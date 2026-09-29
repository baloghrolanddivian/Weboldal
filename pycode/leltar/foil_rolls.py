"""Upload, persistence, calculation, and UI helpers for edge-band roll inventory."""

from __future__ import annotations

import csv
import html
import io
import json
import math
import re
import secrets
import unicodedata
from datetime import datetime
from pathlib import Path

from tools.excel import normalize_excel_payload

from .routes import (
    ADMIN_FOIL_ROLL_INVENTORY_ROUTE,
    ADMIN_INVENTORY_GROUP_ROUTE,
    FOIL_ROLL_INVENTORY_PROCESS_ROUTE,
    FOIL_ROLL_INVENTORY_STATE_ROUTE,
    PRODUCTION_INVENTORY_GROUP_ROUTE,
)

try:
    from openpyxl import load_workbook
except Exception:  # pragma: no cover
    load_workbook = None


ALLOWED_EXTENSIONS = {".xls", ".xlsx", ".xlsm", ".csv"}
_runtime_dir = Path("runtime/folia-tekercs")


def configure_foil_rolls(runtime_dir: Path) -> None:
    """Set the runtime directory used for the active foil inventory session."""
    global _runtime_dir
    _runtime_dir = Path(runtime_dir)


def session_path() -> Path:
    return _runtime_dir / "session.json"


def load_session() -> dict | None:
    path = session_path()
    if not path.is_file():
        return None
    try:
        payload = json.loads(path.read_text(encoding="utf-8"))
    except (OSError, json.JSONDecodeError):
        return None
    return payload if isinstance(payload, dict) else None


def save_session(session: dict) -> None:
    _runtime_dir.mkdir(parents=True, exist_ok=True)
    session_path().write_text(json.dumps(session, ensure_ascii=False, indent=2), encoding="utf-8")


def parse_thickness_mm(description: object) -> float | None:
    """Read the second value from dimensions such as ``21x0,4mm``."""
    text = str(description or "")
    matches = re.findall(r"(\d+(?:[.,]\d+)?)\s*[x×]\s*(\d+(?:[.,]\d+)?)\s*mm\b", text, flags=re.IGNORECASE)
    if not matches:
        return None
    value = _parse_positive(matches[-1][1])
    return value


def calculate_roll_meters(outer_diameter_mm: object, inner_diameter_mm: object, thickness_mm: object) -> float:
    """Estimate wound strip length from annulus area, returning metres."""
    outer = _require_positive(outer_diameter_mm, "A külső átmérő")
    inner = _require_positive(inner_diameter_mm, "A belső átmérő")
    thickness = _require_positive(thickness_mm, "A fóliavastagság")
    if outer <= inner:
        raise ValueError("A külső átmérőnek nagyobbnak kell lennie a belsőnél.")
    return math.pi * (outer * outer - inner * inner) / (4.0 * thickness * 1000.0)


def build_session(file_name: str, payload: bytes) -> dict:
    """Build an editable foil inventory session from an Excel or CSV list."""
    if Path(file_name or "").suffix.lower() not in ALLOWED_EXTENSIONS:
        raise ValueError("A fóliajegyzék csak XLS, XLSX, XLSM vagy CSV lehet.")
    source_rows = _read_rows(file_name, payload)
    rows = []
    for index, source in enumerate(source_rows, start=1):
        part_number = _clean(source.get("part_number"))
        description = _clean(source.get("description"))
        if not part_number and not description:
            continue
        thickness = parse_thickness_mm(description)
        rows.append(
            {
                "row_id": secrets.token_hex(8),
                "part_number": part_number,
                "description": description,
                "book_qty": _clean_number(source.get("book_qty")),
                "thickness_mm": _format_number(thickness) if thickness is not None else "",
                "manual_meters": "",
                "rolls": [],
                "source_index": index,
            }
        )
    if not rows:
        raise ValueError("A feltöltött listában nincs feldolgozható fóliatétel.")
    rows.sort(key=lambda row: (row["description"].casefold(), row["part_number"].casefold()))
    now = datetime.now().isoformat(timespec="seconds")
    return {
        "session_id": secrets.token_hex(6),
        "source_name": Path(file_name).name,
        "created_at": now,
        "updated_at": now,
        "rows": rows,
    }


def apply_action(session: dict, form: dict[str, str]) -> dict:
    """Apply one UI mutation and return the affected row summary."""
    row = _find_row(session, form.get("row_id", ""))
    if row is None:
        raise ValueError("A kiválasztott fóliatétel nem található.")
    action = str(form.get("action", "")).strip().lower()
    if action == "set_value":
        field = str(form.get("field", "")).strip()
        if field not in {"manual_meters", "thickness_mm"}:
            raise ValueError("Nem módosítható mező.")
        raw_value = str(form.get("value", "")).strip()
        if raw_value:
            value = _require_non_negative(raw_value, "A megadott érték")
            if field == "thickness_mm" and value <= 0:
                raise ValueError("A fóliavastagságnak nullánál nagyobbnak kell lennie.")
            row[field] = _format_number(value)
        else:
            row[field] = ""
        if field == "thickness_mm":
            _recalculate_rolls(row)
    elif action == "add_roll":
        thickness = row.get("thickness_mm", "")
        outer = _require_positive(form.get("outer_diameter_mm", ""), "A külső átmérő")
        inner = _require_positive(form.get("inner_diameter_mm", ""), "A belső átmérő")
        meters = calculate_roll_meters(outer, inner, thickness)
        row.setdefault("rolls", []).append(
            {
                "roll_id": secrets.token_hex(5),
                "outer_diameter_mm": _format_number(outer),
                "inner_diameter_mm": _format_number(inner),
                "meters": _format_number(meters, 1),
            }
        )
    elif action == "delete_roll":
        roll_id = str(form.get("roll_id", "")).strip()
        rolls = [roll for roll in row.get("rolls", []) if isinstance(roll, dict)]
        remaining = [roll for roll in rolls if str(roll.get("roll_id", "")) != roll_id]
        if len(remaining) == len(rolls):
            raise ValueError("A kiválasztott tekercsmérés nem található.")
        row["rolls"] = remaining
    else:
        raise ValueError("Ismeretlen művelet.")
    session["updated_at"] = datetime.now().isoformat(timespec="seconds")
    return row_summary(row)


def row_summary(row: dict) -> dict:
    manual = _parse_non_negative(row.get("manual_meters")) or 0.0
    rolls = [roll for roll in row.get("rolls", []) if isinstance(roll, dict)]
    roll_total = sum(_parse_non_negative(roll.get("meters")) or 0.0 for roll in rolls)
    return {
        "row_id": str(row.get("row_id", "")),
        "manual_meters": str(row.get("manual_meters", "")),
        "thickness_mm": str(row.get("thickness_mm", "")),
        "rolls": rolls,
        "roll_total_meters": _format_number(roll_total, 1),
        "total_meters": _format_number(manual + roll_total, 1),
    }


def render_page(admin: bool = False, message: str = "", success: bool = False) -> bytes:
    """Render the admin upload view or the production measurement view."""
    session = load_session()
    back_route = ADMIN_INVENTORY_GROUP_ROUTE if admin else PRODUCTION_INVENTORY_GROUP_ROUTE
    notice = f'<div class="foil-notice{" is-success" if success else ""}">{html.escape(message)}</div>' if message else ""
    upload = ""
    if admin:
        source = html.escape(str(session.get("source_name", ""))) if session else "nincs aktív lista"
        upload = f'''
        <section class="foil-panel foil-upload">
          <div><span class="foil-eyebrow">Lista frissítése</span><h2>Aktuális fóliák feltöltése</h2><p>Excel vagy CSV szükséges, legalább alkatrészszám és leírás oszloppal. Aktív forrás: <strong>{source}</strong></p></div>
          <form method="post" action="{FOIL_ROLL_INVENTORY_PROCESS_ROUTE}" enctype="multipart/form-data">
            <input type="file" name="stock_file" accept=".xls,.xlsx,.xlsm,.csv" required />
            <button type="submit">Fólista betöltése</button>
          </form>
        </section>'''
    if session is None:
        content = '<section class="foil-panel foil-empty"><h2>Még nincs aktív fólista</h2><p>Az admin felületen előbb töltsd fel az aktuális élfóliákat.</p></section>'
    else:
        cards = "".join(_render_row(row) for row in session.get("rows", []) if isinstance(row, dict))
        content = f'''
        <section class="foil-panel foil-workspace" data-foil-root data-state-route="{FOIL_ROLL_INVENTORY_STATE_ROUTE}">
          <div class="foil-heading"><div><span class="foil-eyebrow">{html.escape(str(session.get("source_name", "")))}</span><h2>Tekercsek felmérése</h2><p>A kézzel megadott méterhez automatikusan hozzáadjuk a lemért tekercseket. A becslés a külső és belső átmérő, valamint a módosítható fóliavastagság alapján készül.</p></div><strong>{len(session.get("rows", []))} tétel</strong></div>
          <label class="foil-search"><span>Keresés</span><input type="search" data-foil-search placeholder="Név vagy alkatrészszám…" autocomplete="off" /></label>
          <div class="foil-list">{cards}</div>
        </section>'''
    page = f'''<!doctype html><html lang="hu"><head><meta charset="utf-8"><meta name="viewport" content="width=device-width,initial-scale=1"><title>Divian-HUB | Fóliatekercs-számító</title><link rel="stylesheet" href="/styles.css"><style>{_styles()}</style></head>
    <body class="foil-page" data-theme-scope="default"><main class="foil-shell"><header class="foil-top"><div><span class="foil-eyebrow">Divian-HUB · Leltár</span><h1>Fóliatekercs-számító</h1></div><a href="{back_route}">Vissza a modulokhoz</a></header>{notice}{upload}{content}</main>{_script()}</body></html>'''
    return page.encode("utf-8")


def _render_row(row: dict) -> str:
    summary = row_summary(row)
    row_id = html.escape(str(row.get("row_id", "")), quote=True)
    rolls = "".join(_render_roll(roll) for roll in summary["rolls"])
    if not rolls:
        rolls = '<span class="foil-no-rolls" data-empty-rolls>Még nincs hozzáadott tekercs.</span>'
    return f'''
      <details class="foil-card" data-foil-row data-row-id="{row_id}" data-search="{html.escape((str(row.get('part_number', '')) + ' ' + str(row.get('description', ''))).casefold(), quote=True)}">
        <summary class="foil-card-head"><div class="foil-item-name"><strong>{html.escape(str(row.get("description", "")) or "Név nélküli fólia")}</strong><span>{html.escape(str(row.get("part_number", "")) or "–")} · <b data-thickness-summary>{html.escape(summary['thickness_mm'] or '–')}</b> mm</span></div><div class="foil-card-result"><div class="foil-total"><small>Összesen</small><b data-total>{summary["total_meters"]}</b><em>m</em></div><span class="foil-measure-action">Mérés</span></div></summary>
        <div class="foil-card-body"><div class="foil-entry-row"><div class="foil-fields">
          <label><span>Egész tekercsek</span><div><input data-row-field="manual_meters" inputmode="decimal" value="{html.escape(summary['manual_meters'], quote=True)}" placeholder="pl. 600" /><i>m</i></div></label>
          <label><span>Fóliavastagság</span><div><input data-row-field="thickness_mm" inputmode="decimal" value="{html.escape(summary['thickness_mm'], quote=True)}" placeholder="pl. 0,4" /><i>mm</i></div></label>
        </div>
          <form class="foil-add-roll" data-add-roll><label><span>Külső átmérő</span><div><input name="outer_diameter_mm" inputmode="decimal" placeholder="pl. 380" required /><i>mm</i></div></label><label><span>Belső átmérő</span><div><input name="inner_diameter_mm" inputmode="decimal" placeholder="pl. 90" required /><i>mm</i></div></label><button type="submit">Hozzáadás</button></form>
        </div>
        <div class="foil-roll-section"><div class="foil-roll-title"><strong>Darab tekercsek</strong><span>Tekercsek együtt: <b data-roll-total>{summary['roll_total_meters']}</b> m</span></div><div class="foil-rolls" data-roll-list>{rolls}</div>
          <div class="foil-error" data-row-error role="alert"></div>
        </div></div>
      </details>'''


def _render_roll(roll: dict) -> str:
    roll_id = html.escape(str(roll.get("roll_id", "")), quote=True)
    return f'''<div class="foil-roll" data-roll-id="{roll_id}"><span><b>{html.escape(str(roll.get("meters", "0")))} m</b><small>Ø {html.escape(str(roll.get("outer_diameter_mm", "")))} / {html.escape(str(roll.get("inner_diameter_mm", "")))} mm</small></span><button type="button" data-delete-roll aria-label="Tekercsmérés törlése">×</button></div>'''


def _styles() -> str:
    return """
    :root{--foil:#0f766e;--foil2:#14b8a6;--ink:#0f172a;--muted:#64748b;--line:#dbe4ea}*{box-sizing:border-box}.foil-page{margin:0;background:radial-gradient(circle at 8% 0,#ccfbf1 0,transparent 32%),#f4f7f8;color:var(--ink);font-family:Manrope,Arial,sans-serif}.foil-shell{width:min(1380px,calc(100% - 28px));margin:16px auto 44px;display:grid;gap:16px}.foil-top,.foil-panel{background:rgba(255,255,255,.96);border:1px solid rgba(15,23,42,.08);border-radius:26px;box-shadow:0 22px 55px rgba(15,23,42,.08)}.foil-top{display:flex;justify-content:space-between;align-items:center;padding:21px 24px}.foil-top h1,.foil-panel h2{margin:5px 0 0;font:800 1.55rem/1.1 "Space Grotesk",sans-serif}.foil-top a{color:var(--ink);text-decoration:none;font-weight:900}.foil-eyebrow{color:var(--foil);font-size:.75rem;font-weight:900;text-transform:uppercase;letter-spacing:.08em}.foil-notice{padding:13px 16px;border-radius:16px;background:#fff7ed;color:#9a3412;border:1px solid #fed7aa;font-weight:800}.foil-notice.is-success{background:#ecfdf5;color:#047857;border-color:#bbf7d0}.foil-upload{display:grid;grid-template-columns:1fr auto;gap:22px;align-items:end;padding:22px}.foil-upload p,.foil-heading p,.foil-empty p{margin:7px 0 0;color:var(--muted)}.foil-upload form{display:flex;gap:10px;align-items:center}.foil-upload input{max-width:320px;padding:11px;border:1px solid var(--line);border-radius:14px}.foil-upload button,.foil-add-roll button{min-height:44px;padding:0 16px;border:0;border-radius:14px;background:var(--foil);color:#fff;font-weight:900;cursor:pointer}.foil-workspace{padding:20px}.foil-heading{display:flex;justify-content:space-between;gap:16px}.foil-heading>strong{align-self:start;padding:10px 14px;border-radius:999px;background:#ccfbf1;color:#115e59}.foil-search{display:grid;grid-template-columns:auto minmax(220px,430px);align-items:center;gap:10px;margin:18px 0;color:var(--muted);font-size:.78rem;font-weight:900;text-transform:uppercase}.foil-search input{min-height:44px;padding:0 15px;border:1px solid var(--line);border-radius:999px;font:800 .95rem Manrope;text-transform:none}.foil-list{display:grid;gap:8px}.foil-card{overflow:hidden;border:1px solid var(--line);border-radius:17px;background:#fff}.foil-card-head{display:flex;justify-content:space-between;gap:15px;align-items:center;min-height:64px;padding:10px 12px 10px 16px;background:linear-gradient(135deg,#f8fafc,#ecfdf5);cursor:pointer;list-style:none}.foil-card-head::-webkit-details-marker{display:none}.foil-item-name{min-width:0}.foil-item-name strong{display:block;overflow:hidden;text-overflow:ellipsis;white-space:nowrap;font-size:.94rem}.foil-item-name span{display:block;margin-top:3px;color:var(--muted);font-size:.76rem;font-weight:800}.foil-card-result{display:flex;align-items:center;gap:12px}.foil-total{display:flex;align-items:baseline;gap:4px;white-space:nowrap}.foil-total small{margin-right:3px;color:var(--muted);font-size:.72rem;font-weight:800}.foil-total b{font:800 1.25rem "Space Grotesk"}.foil-total em{color:var(--foil);font-style:normal;font-size:.82rem;font-weight:900}.foil-measure-action{display:inline-flex;align-items:center;justify-content:center;min-height:34px;padding:0 12px;border-radius:11px;background:var(--foil);color:#fff;font-size:.78rem;font-weight:900}.foil-card[open] .foil-measure-action{background:#ccfbf1;color:#115e59}.foil-card[open] .foil-measure-action::after{content:" bezárása"}.foil-card-body{border-top:1px solid #e8eef2}.foil-fields{display:grid;grid-template-columns:1fr 1fr;gap:12px;padding:16px 18px}.foil-fields label,.foil-add-roll label{display:grid;gap:6px;color:#475569;font-size:.78rem;font-weight:900}.foil-fields label>div,.foil-add-roll label>div{display:flex;align-items:center;border:1px solid var(--line);border-radius:14px;background:#fff;overflow:hidden}.foil-fields input,.foil-add-roll input{width:100%;min-width:0;height:43px;padding:0 12px;border:0;outline:0;font:800 1rem Manrope}.foil-fields i,.foil-add-roll i{padding:0 12px;color:var(--muted);font-style:normal}.foil-fields small{color:#94a3b8;font-weight:700}.foil-roll-section{padding:0 18px 18px}.foil-roll-title{display:flex;justify-content:space-between;gap:10px;padding-top:14px;border-top:1px solid #eef2f7}.foil-roll-title span{color:var(--muted);font-size:.84rem}.foil-rolls{display:flex;flex-wrap:wrap;gap:8px;margin:12px 0}.foil-no-rolls{color:#94a3b8;font-size:.84rem}.foil-roll{display:flex;align-items:center;gap:8px;padding:8px 8px 8px 12px;border-radius:14px;background:#ecfdf5;border:1px solid #bbf7d0}.foil-roll span b,.foil-roll span small{display:block}.foil-roll span small{margin-top:2px;color:#64748b;font-size:.72rem}.foil-roll button{width:28px;height:28px;border:0;border-radius:9px;background:#fff;color:#dc2626;font-size:1.2rem;cursor:pointer}.foil-add-roll{display:grid;grid-template-columns:1fr 1fr auto;gap:10px;align-items:end}.foil-error{min-height:18px;margin-top:7px;color:#b91c1c;font-size:.82rem;font-weight:800}.foil-empty{padding:28px}.foil-entry-row{display:grid;grid-template-columns:minmax(100px,1fr) minmax(100px,1fr) minmax(330px,2.25fr);gap:8px;align-items:end;padding:12px}.foil-entry-row>.foil-fields{display:contents}.foil-entry-row>.foil-add-roll{display:grid;grid-template-columns:repeat(2,minmax(100px,1fr)) auto;gap:7px;align-items:end;padding:6px;border:1px solid rgba(15,118,110,.28);border-radius:14px;background:#f0fdfa;box-shadow:inset 0 0 0 1px rgba(255,255,255,.75)}.foil-entry-row .foil-add-roll button{grid-column:auto;width:auto;min-height:43px;white-space:nowrap}.foil-entry-row+.foil-roll-section{padding:0 12px 12px}.foil-entry-row+.foil-roll-section .foil-roll-title{padding-top:10px}@media(max-width:1100px){.foil-shell{width:calc(100% - 16px);margin:8px auto 24px}.foil-top{padding:17px;align-items:flex-start}.foil-top h1{font-size:1.25rem}.foil-upload{grid-template-columns:1fr}.foil-upload form{display:grid}.foil-workspace{padding:12px}.foil-card-head{min-height:58px;padding:9px 10px 9px 13px}.foil-item-name strong{max-width:48vw}.foil-card-result{gap:7px}.foil-total small{display:none}.foil-measure-action{padding:0 9px}.foil-card[open] .foil-measure-action::after{content:""}.foil-fields input,.foil-add-roll input{height:38px;padding:0 7px;font-size:.9rem}.foil-fields i,.foil-add-roll i{padding:0 7px}.foil-fields label,.foil-add-roll label{gap:4px;font-size:.68rem}.foil-roll-title{display:grid}.foil-rolls{margin:8px 0}.foil-entry-row{grid-template-columns:minmax(100px,1fr) minmax(100px,1fr) minmax(310px,2.25fr);overflow-x:auto}.foil-entry-row .foil-add-roll button{min-height:38px}}
    """


def _script() -> str:
    return f'''<script>(()=>{{const root=document.querySelector('[data-foil-root]');if(!root)return;const route=root.dataset.stateRoute;const post=(row,data)=>{{data.set('row_id',row.dataset.rowId||'');return fetch(route,{{method:'POST',headers:{{'Content-Type':'application/x-www-form-urlencoded; charset=UTF-8'}},body:data.toString(),credentials:'same-origin',cache:'no-store'}}).then(async r=>{{const p=await r.json().catch(()=>({{error:'A mentés nem sikerült.'}}));if(!r.ok)throw new Error(p.error||'A mentés nem sikerült.');return p;}})}};const rollHtml=(roll)=>{{const el=document.createElement('div');el.className='foil-roll';el.dataset.rollId=roll.roll_id;el.innerHTML='<span><b></b><small></small></span><button type="button" data-delete-roll aria-label="Tekercsmérés törlése">×</button>';el.querySelector('b').textContent=roll.meters+' m';el.querySelector('small').textContent='Ø '+roll.outer_diameter_mm+' / '+roll.inner_diameter_mm+' mm';return el;}};const paint=(row,p)=>{{row.querySelector('[data-total]').textContent=p.total_meters;row.querySelector('[data-roll-total]').textContent=p.roll_total_meters;row.querySelector('[data-thickness-summary]').textContent=p.thickness_mm||'–';const list=row.querySelector('[data-roll-list]');list.replaceChildren(...p.rolls.map(rollHtml));if(!p.rolls.length){{const empty=document.createElement('span');empty.className='foil-no-rolls';empty.dataset.emptyRolls='';empty.textContent='Még nincs hozzáadott tekercs.';list.append(empty);}}}};root.querySelectorAll('[data-foil-row]').forEach(row=>{{const error=row.querySelector('[data-row-error]');row.querySelectorAll('[data-row-field]').forEach(input=>{{const save=()=>{{const data=new URLSearchParams();data.set('action','set_value');data.set('field',input.dataset.rowField);data.set('value',input.value);post(row,data).then(p=>{{paint(row,p);error.textContent='';}}).catch(e=>error.textContent=e.message);}};input.addEventListener('change',save);input.addEventListener('keydown',e=>{{if(e.key==='Enter'){{e.preventDefault();input.blur();}}}});}});row.querySelector('[data-add-roll]').addEventListener('submit',e=>{{e.preventDefault();const form=e.currentTarget;const data=new URLSearchParams(new FormData(form));data.set('action','add_roll');post(row,data).then(p=>{{paint(row,p);form.reset();error.textContent='';}}).catch(err=>error.textContent=err.message);}});row.querySelector('[data-roll-list]').addEventListener('click',e=>{{const button=e.target.closest('[data-delete-roll]');if(!button)return;const data=new URLSearchParams();data.set('action','delete_roll');data.set('roll_id',button.closest('[data-roll-id]').dataset.rollId||'');post(row,data).then(p=>{{paint(row,p);error.textContent='';}}).catch(err=>error.textContent=err.message);}});}});const search=root.querySelector('[data-foil-search]');search?.addEventListener('input',()=>{{const q=search.value.trim().toLocaleLowerCase('hu-HU');root.querySelectorAll('[data-foil-row]').forEach(row=>row.hidden=q&&!row.dataset.search.includes(q));}});}})();</script>'''


def _recalculate_rolls(row: dict) -> None:
    thickness = row.get("thickness_mm", "")
    if not str(thickness).strip():
        return
    for roll in row.get("rolls", []):
        if not isinstance(roll, dict):
            continue
        roll["meters"] = _format_number(
            calculate_roll_meters(roll.get("outer_diameter_mm"), roll.get("inner_diameter_mm"), thickness), 1
        )


def _find_row(session: dict, row_id: object) -> dict | None:
    clean_id = str(row_id or "").strip()
    return next((row for row in session.get("rows", []) if isinstance(row, dict) and row.get("row_id") == clean_id), None)


def _read_rows(file_name: str, payload: bytes) -> list[dict]:
    suffix = Path(file_name).suffix.lower()
    if suffix == ".csv":
        text = payload.decode("utf-8-sig", errors="replace")
        dialect = csv.Sniffer().sniff(text[:2048], delimiters=";,\t,")
        raw_rows = list(csv.reader(io.StringIO(text), dialect))
    else:
        if load_workbook is None:
            raise RuntimeError("Az Excel beolvasásához hiányzik az openpyxl csomag.")
        workbook = load_workbook(io.BytesIO(normalize_excel_payload(payload)), read_only=True, data_only=True)
        raw_rows = list(workbook.active.iter_rows(values_only=True))
    if not raw_rows:
        raise ValueError("A feltöltött fólista üres.")
    header_map = _header_map(raw_rows[0])
    missing = [label for key, label in (("part_number", "Alkatr.-szám"), ("description", "Alkatr.-leírás")) if key not in header_map]
    if missing:
        raise ValueError("Hiányzó kötelező oszlop: " + ", ".join(missing))
    result = []
    for values in raw_rows[1:]:
        result.append({key: values[index] if index < len(values) else "" for key, index in header_map.items()})
    return result


def _header_map(headers: tuple | list) -> dict[str, int]:
    aliases = {
        "part_number": {"alkatr.-szam", "alkatresz szam", "alkatr-szam", "cikkszam"},
        "description": {"alkatr.-leiras", "alkatresz leiras", "leiras", "megnevezes"},
        "book_qty": {"konyvelesi mennyiseg", "konyvelt mennyiseg", "raktari keszlet", "keszlet"},
    }
    normalized = [_normalize(value) for value in headers]
    result = {}
    for key, names in aliases.items():
        for index, header in enumerate(normalized):
            if header in names:
                result[key] = index
                break
    return result


def _normalize(value: object) -> str:
    text = unicodedata.normalize("NFKD", _clean(value).casefold())
    return re.sub(r"\s+", " ", "".join(char for char in text if not unicodedata.combining(char))).strip()


def _clean(value: object) -> str:
    return re.sub(r"\s+", " ", str(value or "")).strip()


def _clean_number(value: object) -> str:
    parsed = _parse_non_negative(value)
    return _format_number(parsed) if parsed is not None else _clean(value)


def _parse_positive(value: object) -> float | None:
    parsed = _parse_non_negative(value)
    return parsed if parsed is not None and parsed > 0 else None


def _parse_non_negative(value: object) -> float | None:
    text = str(value or "").strip().replace(" ", "").replace(",", ".")
    if not text:
        return None
    try:
        parsed = float(text)
    except ValueError:
        return None
    return parsed if math.isfinite(parsed) and parsed >= 0 else None


def _require_positive(value: object, label: str) -> float:
    parsed = _parse_positive(value)
    if parsed is None:
        raise ValueError(f"{label} csak nullánál nagyobb szám lehet.")
    return parsed


def _require_non_negative(value: object, label: str) -> float:
    parsed = _parse_non_negative(value)
    if parsed is None:
        raise ValueError(f"{label} csak nem negatív szám lehet.")
    return parsed


def _format_number(value: float | None, decimals: int = 3) -> str:
    if value is None:
        return ""
    text = f"{value:.{decimals}f}".rstrip("0").rstrip(".")
    return text.replace(".", ",")
