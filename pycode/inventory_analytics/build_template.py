"""Build the downloadable, human-friendly inventory analytics input template."""

from __future__ import annotations

from datetime import date, timedelta
from pathlib import Path

from openpyxl import Workbook
from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation


SHEETS = {
    "Termékek": [
        "Cikkszám", "Megnevezés", "Kategória", "Mértékegység", "Beszerzési egységár",
        "Átfutási idő (nap)", "Biztonsági készlet",
    ],
    "Készlet": ["Dátum", "Raktár", "Cikkszám", "Készlet", "Foglalt készlet"],
    "Mozgások": ["Mozgás azonosító", "Dátum", "Raktár", "Cikkszám", "Mozgástípus", "Mennyiség"],
    "Nyitott rendelések": ["Rendeléssor azonosító", "Várható dátum", "Raktár", "Cikkszám", "Nyitott mennyiség"],
}


def build_template(output_path: Path) -> None:
    workbook = Workbook()
    workbook.remove(workbook.active)
    guide = workbook.create_sheet("Útmutató")
    for sheet_name in SHEETS:
        workbook.create_sheet(sheet_name)
    lists = workbook.create_sheet("Listák")

    guide.sheet_view.showGridLines = False
    guide.merge_cells("A1:H2")
    guide["A1"] = "KÉSZLET INTELLIGENCE · ADATFRISSÍTÉSI SABLON"
    guide["A1"].font = Font(name="Aptos Display", size=22, bold=True, color="FFFFFF")
    guide["A1"].fill = PatternFill("solid", fgColor="17221B")
    guide["A1"].alignment = Alignment(vertical="center")
    guide.row_dimensions[1].height = 30
    guide.row_dimensions[2].height = 18
    guide["A4"] = "Kitöltési sorrend"
    guide["A4"].font = Font(size=14, bold=True, color="17221B")
    instructions = [
        ("1. Termékek", "A teljes aktív cikktörzs. Ár, átfutási idő vagy biztonsági készlet változásakor frissítsd."),
        ("2. Készlet", "Minden raktár és cikk aktuális, dátumozott pillanatképe. A nulla készletű aktív cikkeket is add át."),
        ("3. Mozgások", "Az előző feltöltés óta keletkezett mozgások. Az azonosító legyen tartósan egyedi."),
        ("4. Nyitott rendelések", "Az összes jelenleg nyitott beszerzési sor. Minden alkalommal teljes állományként add át."),
    ]
    for row, (title, description) in enumerate(instructions, start=6):
        guide[f"A{row}"] = title
        guide[f"A{row}"].font = Font(bold=True, color="526A11")
        guide[f"B{row}"] = description
        guide[f"B{row}"].alignment = Alignment(wrap_text=True)
    guide["A12"] = "Feltöltési gyakoriság"
    guide["A12"].font = Font(size=14, bold=True)
    guide["A14"] = "Ajánlott"
    guide["B14"] = "Naponta vagy hetente, mindig ugyanazon zárási időpont után."
    guide["A15"] = "Szakmai minimum"
    guide["B15"] = "Heti készlet + folyamatos mozgások + legalább 90 nap történet."
    guide["A16"] = "Szezonalitáshoz"
    guide["B16"] = "12–24 hónap mozgástörténet."
    guide["A18"] = "Fontos szabályok"
    guide["A18"].font = Font(size=14, bold=True)
    for row, text in enumerate([
        "A fejlécneveket és a munkalapneveket ne írd át.",
        "Egy cikkszám minden lapon pontosan ugyanabban a formában szerepeljen.",
        "Negatív kiadási mennyiség is elfogadott; az elemzés abszolút kiadásként kezeli.",
        "A Mozgás azonosító és a Rendeléssor azonosító vállalati szinten legyen egyedi.",
        "Képlet helyett lehetőleg értékeket exportálj.",
    ], start=20):
        guide[f"A{row}"] = "•"
        guide[f"B{row}"] = text
    guide.column_dimensions["A"].width = 24
    guide.column_dimensions["B"].width = 90
    guide.freeze_panes = "A4"

    today = date.today()
    examples = {
        "Termékek": [
            ["MINTA-001", "Minta alapanyag", "Alapanyag", "db", 1250, 14, 20],
            ["MINTA-002", "Minta vasalat", "Vasalat", "db", 480, 30, 50],
        ],
        "Készlet": [
            [today, "Központi raktár", "MINTA-001", 120, 15],
            [today, "Központi raktár", "MINTA-002", 40, 5],
        ],
        "Mozgások": [
            ["MINTA-M-0001", today - timedelta(days=7), "Központi raktár", "MINTA-001", "ANYAGKIADÁS", -18],
            ["MINTA-M-0002", today - timedelta(days=3), "Központi raktár", "MINTA-002", "BEVÉTELEZÉS", 30],
        ],
        "Nyitott rendelések": [
            ["MINTA-R-0001", today + timedelta(days=10), "Központi raktár", "MINTA-002", 100],
        ],
    }
    widths = {
        "Termékek": [18, 34, 22, 16, 22, 22, 22],
        "Készlet": [16, 24, 18, 16, 20],
        "Mozgások": [24, 16, 24, 18, 28, 16],
        "Nyitott rendelések": [26, 18, 24, 18, 22],
    }
    for sheet_name, headers in SHEETS.items():
        sheet = workbook[sheet_name]
        sheet.sheet_view.showGridLines = False
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = f"A1:{_column_letter(len(headers))}5000"
        for column, header in enumerate(headers, start=1):
            cell = sheet.cell(1, column, header)
            cell.fill = PatternFill("solid", fgColor="17221B")
            cell.font = Font(bold=True, color="FFFFFF")
            cell.alignment = Alignment(wrap_text=True, vertical="center")
            sheet.column_dimensions[_column_letter(column)].width = widths[sheet_name][column - 1]
        sheet.row_dimensions[1].height = 34
        for row in examples[sheet_name]:
            sheet.append(row)
        for row in range(2, 5001):
            if sheet_name == "Készlet":
                sheet.cell(row, 1).number_format = "yyyy-mm-dd"
            elif sheet_name == "Mozgások":
                sheet.cell(row, 2).number_format = "yyyy-mm-dd"
            elif sheet_name == "Nyitott rendelések":
                sheet.cell(row, 2).number_format = "yyyy-mm-dd"
        sheet.conditional_formatting.add(
            f"A2:{_column_letter(len(headers))}3",
            CellIsRule(operator="notEqual", formula=['""'], fill=PatternFill("solid", fgColor="F1F7DD")),
        )

    lists.sheet_state = "hidden"
    lists.append(["Mozgástípus"])
    for value in ["BEVÉTELEZÉS", "ELADÁS", "TERMELÉSI_FELHASZNÁLÁS", "ANYAGKIADÁS", "FELHASZNÁLÁS", "SELEJT", "KORREKCIÓ", "ÁTHELYEZÉS"]:
        lists.append([value])
    movement_validation = DataValidation(type="list", formula1="'Listák'!$A$2:$A$9", allow_blank=False)
    workbook["Mozgások"].add_data_validation(movement_validation)
    movement_validation.add("E2:E5000")

    output_path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(output_path)


def _column_letter(number: int) -> str:
    result = ""
    while number:
        number, remainder = divmod(number - 1, 26)
        result = chr(65 + remainder) + result
    return result


if __name__ == "__main__":
    repo_root = Path(__file__).resolve().parents[2]
    build_template(repo_root / "data" / "keszlet-elemzes-sablon.xlsx")
