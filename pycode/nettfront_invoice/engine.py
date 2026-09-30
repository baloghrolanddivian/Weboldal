"""NettFront batch XML parsing, aggregation, and Excel export."""

from __future__ import annotations

import io
import re
import xml.etree.ElementTree as ET
from collections import OrderedDict
from datetime import date

from nettfront.engine import (
    _procurement_code_fallbacks,
    apply_translations,
    load_alkatresz_map,
    load_translations,
)

try:
    from openpyxl import Workbook
except Exception:  # pragma: no cover
    Workbook = None

def _resolve_our_code(base_code: str, parts: dict[str, str]) -> tuple[str, bool]:
    resolved = parts.get(base_code, "")
    if not resolved:
        for fallback in _procurement_code_fallbacks(base_code):
            resolved = parts.get(fallback, "")
            if resolved:
                break
    if not resolved and base_code == "NFAY_ANT_PRAS_357x197":
        resolved = "NFAY_ANT_PRAS_357x197_KA_NO_U_NFA"
    return resolved or base_code, not bool(resolved)


def _as_int(value: object) -> int:
    try:
        return int(float(str(value).replace(",", ".")))
    except (TypeError, ValueError):
        return 0


def parse_batch_xml(xml_data: bytes) -> dict:
    """Parse a NettFront batch XML and translate each line to our base code."""
    if not xml_data or len(xml_data) > 10 * 1024 * 1024:
        raise ValueError("Az XML hiányzik vagy túl nagy.")
    try:
        root = ET.fromstring(xml_data)
    except ET.ParseError as exc:
        raise ValueError(f"Az XML nem olvasható: {exc}") from exc
    if root.tag != "Invoice":
        raise ValueError("Az XML gyökéreleme nem Invoice.")
    invoice_number = str(root.get("InvoiceNo", "")).strip()
    if not invoice_number:
        raise ValueError("Az XML nem tartalmaz InvoiceNo azonosítót.")
    parsed_lines: list[dict] = []
    seen_icns: set[str] = set()
    for order in root.findall("Order"):
        batch_id = str(order.get("OrderNo", "")).strip()
        for line in order.findall("Line"):
            description = str(line.get("Description", "")).strip()
            parts = [part.strip() for part in re.split(r"\s+-\s+", description) if part.strip()]
            size_index = next((index for index, part in enumerate(parts) if re.fullmatch(r"\d+\s*[xX]\s*\d+", part)), -1)
            size = f"{line.get('MAG', '').strip()}x{line.get('SZEL', '').strip()}".strip("x")
            product = parts[0] if size_index > 0 else str(line.get("pdCode", "")).strip()
            color = " - ".join(parts[size_index + 1 :]) if size_index >= 0 else str(line.get("SZIN", "")).strip()
            if size_index >= 0:
                size = parts[size_index].replace(" ", "").lower().replace("x", "x")
            if str(line.get("UVEG", "")).strip().upper() == "I" and "üveges" not in product.lower():
                product = f"{product} - Üveges"
            icns = [str(item.get("icnID", "")).strip() for item in line.findall("Item")]
            icns = [icn for icn in icns if icn]
            duplicate = next((icn for icn in icns if icn in seen_icns), "")
            if duplicate:
                raise ValueError(f"Az ICN többször szerepel az XML-ben: {duplicate}")
            seen_icns.update(icns)
            translated = apply_translations(
                [{"termek": product, "szin": color, "meret": size, "m2": "", "db": str(_as_int(line.get("Qty"))), "ossz_m2": "", "egyseg_ar": "", "netto_ar": ""}],
                load_translations(),
            )[0]
            parsed_lines.append(
                {
                    "batch_id": batch_id,
                    "line_no": str(line.get("invLineNo", "")).strip(),
                    "description": description,
                    "termek": translated.get("termek", product),
                    "szin": translated.get("szin", color),
                    "meret": translated.get("meret", size),
                    "kod": translated.get("kod", ""),
                    "quantity": _as_int(line.get("Qty")),
                    "icns": icns,
                    "xml_attributes": dict(line.attrib),
                }
            )
    if not parsed_lines:
        raise ValueError("Az XML nem tartalmaz feldolgozható Line elemet.")
    return {"invoice_number": invoice_number, "lines": parsed_lines, "icn_count": len(seen_icns)}


def build_invoice_from_batch(batch_data: dict, invoice_code: str = "", invoice_date: str = "") -> dict:
    """Build the persisted invoice schema directly from parsed batch XML."""
    xml_number = str(batch_data.get("invoice_number", "")).strip()
    clean_code = re.sub(r"\s+", "", str(invoice_code or "")).upper()
    if not clean_code:
        clean_code = f"NF{date.today().year}-{xml_number}"
    if not re.fullmatch(r"[A-Z0-9][A-Z0-9._/-]*", clean_code):
        raise ValueError("A számlakód formátuma érvénytelen.")
    if not re.sub(r"\D", "", clean_code).endswith(re.sub(r"\D", "", xml_number)):
        raise ValueError(f"A megadott számlakód ({clean_code}) nem egyezik az XML InvoiceNo értékével ({xml_number}).")
    clean_date = str(invoice_date or "").strip()
    try:
        parsed_date = date.fromisoformat(clean_date)
    except ValueError:
        raise ValueError("A számla dátuma kötelező, formátuma ÉÉÉÉ-HH-NN.")
    if parsed_date.year < 2000:
        raise ValueError("A számla dátuma érvénytelen.")
    parts = load_alkatresz_map()
    source_items: list[dict] = []
    aggregated: OrderedDict[tuple[str, str], dict] = OrderedDict()
    for source_index, xml_line in enumerate(batch_data.get("lines", []), start=1):
        row = dict(xml_line)
        row["source_index"] = str(source_index)
        full_code, missing = _resolve_our_code(str(row.get("kod", "")), parts)
        row.update(
            {
                "our_code": full_code,
                "our_description": "",  # TODO: external item-description API
                "missing_in_our_system": missing,
                "expected_qty": int(row.get("quantity", 0) or 0),
                "counted_qty": 0,
                "read_at": "",
                "read_by": "",
            }
        )
        source_items.append(dict(row))
        key = (str(row.get("batch_id", "")), str(row.get("kod", "")))
        if key not in aggregated:
            aggregated[key] = dict(row)
            aggregated[key]["source_indices"] = [row["source_index"]]
            aggregated[key]["xml_line_nos"] = [row.get("line_no", "")]
            aggregated[key]["available_icns"] = list(row.get("icns", []))
        else:
            item = aggregated[key]
            item["expected_qty"] += row["expected_qty"]
            item["quantity"] = item["expected_qty"]
            item["source_indices"].append(row["source_index"])
            item["xml_line_nos"].append(row.get("line_no", ""))
            item["available_icns"].extend(row.get("icns", []))
    items = list(aggregated.values())
    icn_index: dict[str, int] = {}
    for item_index, item in enumerate(items):
        item["xml_quantity"] = int(item.get("expected_qty", 0) or 0)
        item["xml_icn_count"] = len(item.get("available_icns", []))
        for icn in item.get("available_icns", []):
            if icn in icn_index:
                raise ValueError(f"Az ICN több XML-sorhoz tartozik: {icn}")
            icn_index[icn] = item_index
    missing_codes = sorted({str(item.get("our_code", "")) for item in items if item.get("missing_in_our_system")})
    batch_ids = list(dict.fromkeys(str(item.get("batch_id", "")) for item in items if item.get("batch_id")))
    return {
        "invoice_code": clean_code,
        "invoice_date": clean_date,
        "items": items,
        "source_items": source_items,
        "missing_codes": missing_codes,
        "batch_ids": batch_ids,
        "icn_index": icn_index,
        "batch_xml": {
            "invoice_number": xml_number,
            "line_count": len(batch_data.get("lines", [])),
            "matched_line_count": len(batch_data.get("lines", [])),
            "unmatched_lines": [],
            "icn_count": len(icn_index),
        },
    }


def build_invoice_workbook(record: dict) -> bytes:
    if Workbook is None:
        raise RuntimeError("Az Excel exporthoz az openpyxl csomag szükséges.")
    workbook = Workbook()
    summary = workbook.active
    summary.title = "Összesített tételek"
    headers = [
        "Számlakód", "Batch azonosító", "ICN", "Modell", "Méret", "Szín",
        "Várt darabszám", "Számolt darabszám", "Eltérés", "Saját kód",
        "Saját leírás", "Beolvasás ideje", "Beolvasó", "Státusz", "Lezáró", "Lezárás ideje",
    ]
    summary.append(headers)
    for item in record.get("items", []):
        expected = int(item.get("expected_qty", 0) or 0)
        counted = int(item.get("counted_qty", 0) or 0)
        read_icns = ", ".join(str(entry.get("icn", "")) for entry in item.get("icns_read", []) if isinstance(entry, dict))
        summary.append([
            record.get("invoice_code", ""), item.get("batch_id", ""), read_icns or item.get("icn", ""),
            item.get("termek", ""), item.get("meret", ""), item.get("szin", ""), expected,
            counted, counted - expected, item.get("our_code", ""), item.get("our_description", ""),
            item.get("read_at", ""), item.get("read_by", ""), record.get("status", ""),
            record.get("closed_by", ""), record.get("closed_at", ""),
        ])
    history = workbook.create_sheet("Teljes beolvasási előzmény")
    history.append(["Idő", "Felhasználó", "ICN", "Darabszám", "Eredmény", "Üzenet"])
    for entry in record.get("scan_history", []):
        history.append([entry.get(key, "") for key in ("time", "user", "icn", "quantity", "result", "message")])
    source = workbook.create_sheet("Forrássorok")
    source.append(["Számlakód", "Batch", "Forrássor", "Termék", "Szín", "Méret", "Darabszám", "Saját kód", "Hiányzik"])
    for item in record.get("source_items", []):
        source.append([
            record.get("invoice_code", ""), item.get("batch_id", ""), item.get("source_index", ""),
            item.get("termek", ""), item.get("szin", ""), item.get("meret", ""), item.get("db", ""),
            item.get("our_code", ""), "igen" if item.get("missing_in_our_system") else "nem",
        ])
    for sheet in workbook.worksheets:
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions
        for column in sheet.columns:
            letter = column[0].column_letter
            sheet.column_dimensions[letter].width = min(42, max(12, max(len(str(cell.value or "")) for cell in column) + 2))
    output = io.BytesIO()
    workbook.save(output)
    return output.getvalue()
