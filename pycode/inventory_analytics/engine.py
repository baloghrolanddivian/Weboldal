"""Persistent inventory turnover analytics built from periodic Excel uploads."""

from __future__ import annotations

import hashlib
import io
import math
import re
import sqlite3
import statistics
import unicodedata
from collections import defaultdict
from datetime import date, datetime, timedelta
from pathlib import Path

from openpyxl import load_workbook


_runtime_dir: Path | None = None
OUTFLOW_TYPES = {"ELADAS", "TERMELÉSI_FELHASZNALAS", "ANYAGKIADAS", "FELHASZNALAS", "SELEJT"}

SHEETS = {
    "products": ("Termékek", ("Cikkszám", "Megnevezés", "Kategória", "Mértékegység", "Beszerzési egységár", "Átfutási idő (nap)", "Biztonsági készlet")),
    "stock": ("Készlet", ("Dátum", "Raktár", "Cikkszám", "Készlet", "Foglalt készlet")),
    "movements": ("Mozgások", ("Mozgás azonosító", "Dátum", "Raktár", "Cikkszám", "Mozgástípus", "Mennyiség")),
    "orders": ("Nyitott rendelések", ("Rendeléssor azonosító", "Várható dátum", "Raktár", "Cikkszám", "Nyitott mennyiség")),
}


def configure_inventory_analytics(runtime_dir: Path) -> None:
    global _runtime_dir
    _runtime_dir = Path(runtime_dir)
    _runtime_dir.mkdir(parents=True, exist_ok=True)
    _ensure_database()


def _db_path() -> Path:
    if _runtime_dir is None:
        raise RuntimeError("A készletelemzés nincs konfigurálva.")
    return _runtime_dir / "inventory-analytics.db"


def _connect() -> sqlite3.Connection:
    conn = sqlite3.connect(_db_path())
    conn.row_factory = sqlite3.Row
    return conn


def _ensure_database() -> None:
    with _connect() as conn:
        conn.executescript(
            """
            PRAGMA journal_mode=WAL;
            CREATE TABLE IF NOT EXISTS products (
                sku TEXT PRIMARY KEY,
                name TEXT NOT NULL,
                category TEXT NOT NULL,
                unit TEXT NOT NULL,
                unit_cost REAL NOT NULL,
                lead_time_days INTEGER NOT NULL,
                safety_stock REAL NOT NULL,
                updated_at TEXT NOT NULL
            );
            CREATE TABLE IF NOT EXISTS stock_snapshots (
                snapshot_date TEXT NOT NULL,
                warehouse TEXT NOT NULL,
                sku TEXT NOT NULL,
                on_hand REAL NOT NULL,
                reserved REAL NOT NULL,
                upload_id TEXT NOT NULL,
                PRIMARY KEY (snapshot_date, warehouse, sku)
            );
            CREATE TABLE IF NOT EXISTS movements (
                movement_id TEXT PRIMARY KEY,
                movement_date TEXT NOT NULL,
                warehouse TEXT NOT NULL,
                sku TEXT NOT NULL,
                movement_type TEXT NOT NULL,
                quantity REAL NOT NULL,
                upload_id TEXT NOT NULL
            );
            CREATE TABLE IF NOT EXISTS open_orders (
                order_line_id TEXT PRIMARY KEY,
                expected_date TEXT NOT NULL,
                warehouse TEXT NOT NULL,
                sku TEXT NOT NULL,
                open_quantity REAL NOT NULL,
                upload_id TEXT NOT NULL
            );
            CREATE TABLE IF NOT EXISTS uploads (
                upload_id TEXT PRIMARY KEY,
                file_name TEXT NOT NULL,
                uploaded_at TEXT NOT NULL,
                products_count INTEGER NOT NULL,
                stock_count INTEGER NOT NULL,
                movements_count INTEGER NOT NULL,
                orders_count INTEGER NOT NULL
            );
            CREATE TABLE IF NOT EXISTS erp_stock (
                sku TEXT PRIMARY KEY,
                category TEXT NOT NULL,
                description TEXT NOT NULL,
                physical_stock REAL NOT NULL,
                available_stock REAL NOT NULL,
                open_demand REAL NOT NULL,
                confirmed_demand REAL NOT NULL,
                open_purchase REAL NOT NULL,
                safety_stock REAL NOT NULL,
                storage_capacity REAL NOT NULL,
                length_value TEXT NOT NULL,
                width_value TEXT NOT NULL,
                color TEXT NOT NULL,
                inventory_flag TEXT NOT NULL,
                uploaded_at TEXT NOT NULL,
                upload_id TEXT NOT NULL
            );
            CREATE TABLE IF NOT EXISTS erp_consumption (
                upload_id TEXT NOT NULL,
                uploaded_at TEXT NOT NULL,
                window_days INTEGER NOT NULL,
                sku TEXT NOT NULL,
                category TEXT NOT NULL,
                description TEXT NOT NULL,
                monthly_average REAL NOT NULL,
                report_stock REAL NOT NULL,
                unit TEXT NOT NULL,
                safety_stock REAL NOT NULL,
                inventory_flag TEXT NOT NULL,
                long_description TEXT NOT NULL,
                PRIMARY KEY (upload_id, sku)
            );
            CREATE TABLE IF NOT EXISTS erp_reports (
                upload_id TEXT PRIMARY KEY,
                uploaded_at TEXT NOT NULL,
                file_name TEXT NOT NULL,
                report_type TEXT NOT NULL,
                window_days INTEGER NOT NULL,
                row_count INTEGER NOT NULL
            );
            CREATE TABLE IF NOT EXISTS erp_item_settings (
                sku TEXT PRIMARY KEY,
                discontinued INTEGER NOT NULL DEFAULT 0,
                updated_at TEXT NOT NULL
            );
            CREATE TABLE IF NOT EXISTS erp_item_exclusions (
                sku TEXT PRIMARY KEY,
                reason TEXT NOT NULL,
                updated_at TEXT NOT NULL
            );
            INSERT OR IGNORE INTO erp_item_exclusions (sku, reason, updated_at)
            VALUES ('NPANT_RUCH', 'Nem beszerzési cikk', CURRENT_TIMESTAMP);
            CREATE TABLE IF NOT EXISTS erp_door_stock (
                sku TEXT PRIMARY KEY, category TEXT NOT NULL, description TEXT NOT NULL,
                physical_stock REAL NOT NULL, available_stock REAL NOT NULL,
                open_demand REAL NOT NULL, confirmed_demand REAL NOT NULL,
                open_purchase REAL NOT NULL, safety_stock REAL NOT NULL,
                storage_capacity REAL NOT NULL, length_value TEXT NOT NULL,
                width_value TEXT NOT NULL, color TEXT NOT NULL, inventory_flag TEXT NOT NULL,
                uploaded_at TEXT NOT NULL, upload_id TEXT NOT NULL
            );
            CREATE TABLE IF NOT EXISTS erp_door_consumption (
                upload_id TEXT NOT NULL, uploaded_at TEXT NOT NULL, window_days INTEGER NOT NULL,
                sku TEXT NOT NULL, category TEXT NOT NULL, description TEXT NOT NULL,
                monthly_average REAL NOT NULL, report_stock REAL NOT NULL, unit TEXT NOT NULL,
                safety_stock REAL NOT NULL, inventory_flag TEXT NOT NULL, long_description TEXT NOT NULL,
                PRIMARY KEY (upload_id, sku)
            );
            CREATE TABLE IF NOT EXISTS erp_door_reports (
                upload_id TEXT PRIMARY KEY, uploaded_at TEXT NOT NULL, file_name TEXT NOT NULL,
                report_type TEXT NOT NULL, window_days INTEGER NOT NULL, row_count INTEGER NOT NULL
            );
            CREATE TABLE IF NOT EXISTS erp_door_item_settings (
                sku TEXT PRIMARY KEY, discontinued INTEGER NOT NULL DEFAULT 0, updated_at TEXT NOT NULL
            );
            CREATE TABLE IF NOT EXISTS erp_door_item_exclusions (
                sku TEXT PRIMARY KEY, reason TEXT NOT NULL, updated_at TEXT NOT NULL
            );
            """
        )


def _erp_table_names(dataset: str = "parts") -> dict[str, str]:
    prefix = "erp_door_" if dataset == "doors" else "erp_"
    return {
        "stock": f"{prefix}stock",
        "consumption": f"{prefix}consumption",
        "reports": f"{prefix}reports",
        "settings": f"{prefix}item_settings",
        "exclusions": f"{prefix}item_exclusions",
    }


def _canonical_category(value: str) -> str:
    """Merge ERP category aliases that differ only by legacy naming conventions."""
    category = str(value or "").strip() or "Nincs kategória"
    normalized = "".join(
        character for character in unicodedata.normalize("NFKD", category).upper()
        if not unicodedata.combining(character) and character.isalnum()
    )
    aliases = {
        "CSAVAR": "Csavar",
        "CSOMAGOLOANY": "Csomagoló anyag",
        "DEKORANYAG": "Dekoranyag",
        "EGYEBANYAG": "Egyéb-anyag",
        "ELANYAG": "Élanyagok",
        "ELANYAGOK": "Élanyagok",
        "FOGANTYU": "Fogantyú",
        "KOSPAN": "Kospán",
        "LAPANYAG": "Lapanyagok",
        "LAPANYAGOK": "Lapanyagok",
        "MOSOGATO": "Mosogató tálca",
        "MOSOGATOTALCA": "Mosogató tálca",
        "MUNKALAP": "Munkalap",
        "PANT": "Pánt",
        "UVEG": "Üveg",
        "VASALAT": "Vasalat-Szerelvény",
        "VASALATSZERELVENY": "Vasalat-Szerelvény",
        "VIZZARO": "Vízzáró",
    }
    return aliases.get(normalized, category)


def set_discontinued(sku: str, discontinued: bool, dataset: str = "parts") -> None:
    """Persist a manual discontinued flag independently of ERP refreshes."""
    normalized_sku = str(sku or "").strip()
    if not normalized_sku:
        raise ValueError("Hiányzik a cikkszám.")
    tables = _erp_table_names(dataset)
    with _connect() as conn:
        if conn.execute(f"SELECT 1 FROM {tables['stock']} WHERE sku=?", (normalized_sku,)).fetchone() is None:
            raise ValueError("A cikkszám nem található.")
        conn.execute(
            f"""
            INSERT INTO {tables['settings']} (sku, discontinued, updated_at) VALUES (?, ?, ?)
            ON CONFLICT(sku) DO UPDATE SET discontinued=excluded.discontinued, updated_at=excluded.updated_at
            """,
            (normalized_sku, int(bool(discontinued)), datetime.now().isoformat(timespec="seconds")),
        )


def import_workbook(file_name: str, payload: bytes) -> dict:
    """Validate and persist one standard inventory workbook."""
    if not file_name.lower().endswith((".xlsx", ".xlsm")):
        raise ValueError("Csak XLSX vagy XLSM munkafüzet tölthető fel.")
    if len(payload) > 30 * 1024 * 1024:
        raise ValueError("A munkafüzet legfeljebb 30 MB lehet.")
    workbook = load_workbook(io.BytesIO(payload), read_only=True, data_only=True)
    parsed: dict[str, list[dict]] = {}
    warnings: list[str] = []
    orders_provided = SHEETS["orders"][0] in workbook.sheetnames
    for key, (sheet_name, headers) in SHEETS.items():
        if sheet_name not in workbook.sheetnames:
            if key == "orders":
                parsed[key] = []
                warnings.append("A „Nyitott rendelések” lap hiányzik; rendelésérkezéssel nem számolunk.")
                continue
            raise ValueError(f"Hiányzik a kötelező „{sheet_name}” munkalap.")
        parsed[key] = _read_sheet(workbook[sheet_name], headers, sheet_name)

    if not parsed["products"] or not parsed["stock"]:
        raise ValueError("A Termékek és a Készlet munkalapon legalább egy adatsor szükséges.")

    upload_id = hashlib.sha256(payload).hexdigest()
    uploaded_at = datetime.now().isoformat(timespec="seconds")
    with _connect() as conn:
        already = conn.execute("SELECT 1 FROM uploads WHERE upload_id = ?", (upload_id,)).fetchone() is not None
        if already:
            return {"duplicate": True, "warnings": ["Ezt a fájlt korábban már feldolgoztuk; az adatok nem duplázódtak."]}

        for row_number, row in enumerate(parsed["products"], start=2):
            sku = _required_text(row, "Cikkszám", "Termékek", row_number)
            conn.execute(
                """
                INSERT INTO products (sku, name, category, unit, unit_cost, lead_time_days, safety_stock, updated_at)
                VALUES (?, ?, ?, ?, ?, ?, ?, ?)
                ON CONFLICT(sku) DO UPDATE SET name=excluded.name, category=excluded.category,
                    unit=excluded.unit, unit_cost=excluded.unit_cost, lead_time_days=excluded.lead_time_days,
                    safety_stock=excluded.safety_stock, updated_at=excluded.updated_at
                """,
                (
                    sku,
                    _required_text(row, "Megnevezés", "Termékek", row_number),
                    _text(row["Kategória"]) or "Nincs kategória",
                    _text(row["Mértékegység"]) or "db",
                    _number(row["Beszerzési egységár"], "Termékek", row_number, "Beszerzési egységár", minimum=0),
                    int(_number(row["Átfutási idő (nap)"], "Termékek", row_number, "Átfutási idő (nap)", minimum=0)),
                    _number(row["Biztonsági készlet"], "Termékek", row_number, "Biztonsági készlet", minimum=0),
                    uploaded_at,
                ),
            )

        for row_number, row in enumerate(parsed["stock"], start=2):
            conn.execute(
                """
                INSERT INTO stock_snapshots (snapshot_date, warehouse, sku, on_hand, reserved, upload_id)
                VALUES (?, ?, ?, ?, ?, ?)
                ON CONFLICT(snapshot_date, warehouse, sku) DO UPDATE SET
                    on_hand=excluded.on_hand, reserved=excluded.reserved, upload_id=excluded.upload_id
                """,
                (
                    _date_text(row["Dátum"], "Készlet", row_number),
                    _required_text(row, "Raktár", "Készlet", row_number),
                    _required_text(row, "Cikkszám", "Készlet", row_number),
                    _number(row["Készlet"], "Készlet", row_number, "Készlet"),
                    _number(row["Foglalt készlet"], "Készlet", row_number, "Foglalt készlet", minimum=0),
                    upload_id,
                ),
            )

        inserted_movements = 0
        for row_number, row in enumerate(parsed["movements"], start=2):
            cursor = conn.execute(
                """
                INSERT OR IGNORE INTO movements
                    (movement_id, movement_date, warehouse, sku, movement_type, quantity, upload_id)
                VALUES (?, ?, ?, ?, ?, ?, ?)
                """,
                (
                    _required_text(row, "Mozgás azonosító", "Mozgások", row_number),
                    _date_text(row["Dátum"], "Mozgások", row_number),
                    _required_text(row, "Raktár", "Mozgások", row_number),
                    _required_text(row, "Cikkszám", "Mozgások", row_number),
                    _movement_type(row["Mozgástípus"]),
                    _number(row["Mennyiség"], "Mozgások", row_number, "Mennyiség"),
                    upload_id,
                ),
            )
            inserted_movements += cursor.rowcount

        if orders_provided:
            # This sheet is a full current-state export, not an incremental history.
            conn.execute("DELETE FROM open_orders")
        for row_number, row in enumerate(parsed["orders"], start=2):
            conn.execute(
                """
                INSERT INTO open_orders (order_line_id, expected_date, warehouse, sku, open_quantity, upload_id)
                VALUES (?, ?, ?, ?, ?, ?)
                ON CONFLICT(order_line_id) DO UPDATE SET expected_date=excluded.expected_date,
                    warehouse=excluded.warehouse, sku=excluded.sku,
                    open_quantity=excluded.open_quantity, upload_id=excluded.upload_id
                """,
                (
                    _required_text(row, "Rendeléssor azonosító", "Nyitott rendelések", row_number),
                    _date_text(row["Várható dátum"], "Nyitott rendelések", row_number),
                    _required_text(row, "Raktár", "Nyitott rendelések", row_number),
                    _required_text(row, "Cikkszám", "Nyitott rendelések", row_number),
                    _number(row["Nyitott mennyiség"], "Nyitott rendelések", row_number, "Nyitott mennyiség", minimum=0),
                    upload_id,
                ),
            )

        conn.execute(
            "INSERT INTO uploads VALUES (?, ?, ?, ?, ?, ?, ?)",
            (upload_id, Path(file_name).name, uploaded_at, len(parsed["products"]), len(parsed["stock"]), inserted_movements, len(parsed["orders"])),
        )
        conn.commit()

    if inserted_movements < len(parsed["movements"]):
        warnings.append(f"{len(parsed['movements']) - inserted_movements} már ismert mozgássor kimaradt.")
    return {
        "duplicate": False,
        "products": len(parsed["products"]),
        "stock": len(parsed["stock"]),
        "movements": inserted_movements,
        "orders": len(parsed["orders"]),
        "warnings": warnings,
    }


def dashboard(period_days: int = 90, dataset: str = "parts") -> dict:
    """Return executive KPIs and SKU-level stock diagnostics."""
    period_days = period_days if period_days in {7, 30, 90, 180, 365} else 90
    tables = _erp_table_names(dataset)
    with _connect() as conn:
        has_erp_stock = conn.execute(f"SELECT 1 FROM {tables['stock']} LIMIT 1").fetchone() is not None
    if has_erp_stock:
        return erp_dashboard(period_days, dataset)
    if dataset == "doors":
        return {
            "mode": "erp",
            "dataset": "doors",
            "period_days": period_days,
            "selected_window": 90,
            "latest_snapshot": "",
            "reports": [],
            "available_windows": [],
            "rows": [],
            "kpis": {
                "sku_count": 0, "moving_count": 0, "shortage_count": 0, "negative_count": 0,
                "erp_reconcile_count": 0, "dead_count": 0, "overstock_count": 0,
                "median_turnover": None, "median_turnover_90": None, "median_turnover_365": None,
                "median_turnover_months_90": None, "median_turnover_months_365": None,
                "median_coverage": None, "incoming_count": 0, "discontinued_count": 0,
            },
        }
    today = date.today()
    since = (today - timedelta(days=period_days - 1)).isoformat()
    with _connect() as conn:
        products = {row["sku"]: dict(row) for row in conn.execute("SELECT * FROM products")}
        uploads = [dict(row) for row in conn.execute("SELECT * FROM uploads ORDER BY uploaded_at DESC LIMIT 8")]
        latest_date_row = conn.execute("SELECT MAX(snapshot_date) AS value FROM stock_snapshots").fetchone()
        latest_date = str(latest_date_row["value"] or "")
        current_rows = list(conn.execute(
            "SELECT sku, SUM(on_hand) on_hand, SUM(reserved) reserved FROM stock_snapshots WHERE snapshot_date=? GROUP BY sku",
            (latest_date,),
        )) if latest_date else []
        snapshot_rows = list(conn.execute(
            "SELECT sku, snapshot_date, SUM(on_hand) qty FROM stock_snapshots WHERE snapshot_date>=? GROUP BY sku, snapshot_date",
            (since,),
        ))
        movement_rows = list(conn.execute(
            "SELECT sku, movement_date, movement_type, quantity FROM movements WHERE movement_date>=?",
            (since,),
        ))
        order_rows = list(conn.execute(
            "SELECT sku, SUM(open_quantity) qty, MIN(expected_date) expected FROM open_orders GROUP BY sku"
        ))

    current = {row["sku"]: (float(row["on_hand"]), float(row["reserved"])) for row in current_rows}
    snapshot_qty: dict[str, list[float]] = defaultdict(list)
    for row in snapshot_rows:
        snapshot_qty[row["sku"]].append(float(row["qty"]))
    outbound: dict[str, float] = defaultdict(float)
    weekly: dict[str, dict[str, float]] = defaultdict(lambda: defaultdict(float))
    last_movement: dict[str, str] = {}
    for row in movement_rows:
        sku = row["sku"]
        movement_date = str(row["movement_date"])
        last_movement[sku] = max(last_movement.get(sku, ""), movement_date)
        if row["movement_type"] in OUTFLOW_TYPES:
            qty = abs(float(row["quantity"]))
            outbound[sku] += qty
            day = date.fromisoformat(movement_date)
            week_key = f"{day.isocalendar().year}-{day.isocalendar().week:02d}"
            weekly[sku][week_key] += qty
    orders = {row["sku"]: (float(row["qty"]), str(row["expected"] or "")) for row in order_rows}

    total_outbound_value = sum(outbound[sku] * float(product["unit_cost"]) for sku, product in products.items())
    ranked_skus = sorted(products, key=lambda sku: outbound[sku] * float(products[sku]["unit_cost"]), reverse=True)
    abc: dict[str, str] = {}
    cumulative = 0.0
    for sku in ranked_skus:
        share_before = cumulative / total_outbound_value if total_outbound_value else 0
        cumulative += outbound[sku] * float(products[sku]["unit_cost"])
        abc[sku] = "A" if share_before < 0.80 else ("B" if share_before < 0.95 else "C")

    rows = []
    for sku, product in products.items():
        on_hand, reserved = current.get(sku, (0.0, 0.0))
        available = on_hand - reserved
        avg_stock = statistics.fmean(snapshot_qty[sku]) if snapshot_qty[sku] else on_hand
        used = outbound[sku]
        daily_use = used / period_days
        coverage = available / daily_use if daily_use > 0 else None
        turnover = used / avg_stock if avg_stock > 0 else None
        lead_time = int(product["lead_time_days"])
        reorder_point = daily_use * lead_time + float(product["safety_stock"])
        incoming, expected = orders.get(sku, (0.0, ""))
        week_values = list(weekly[sku].values())
        if len(week_values) >= 2 and statistics.fmean(week_values) > 0:
            cv = statistics.pstdev(week_values) / statistics.fmean(week_values)
            xyz = "X" if cv <= 0.5 else ("Y" if cv <= 1.0 else "Z")
        else:
            xyz = "Z" if used > 0 else "-"
        days_since_move = (today - date.fromisoformat(last_movement[sku])).days if sku in last_movement else None
        if available <= 0 and daily_use > 0:
            status, severity = "Készlethiány", 4
        elif available + incoming < reorder_point:
            status, severity = "Utánrendelés szükséges", 3
        elif used == 0 and on_hand > 0:
            status, severity = "Nem forgó készlet", 2
        elif coverage is not None and coverage > 120:
            status, severity = "Túlkészlet", 1
        else:
            status, severity = "Rendben", 0
        suggested = max(0.0, reorder_point + daily_use * 30 - available - incoming)
        rows.append({
            "sku": sku, "name": product["name"], "category": product["category"], "unit": product["unit"],
            "on_hand": on_hand, "available": available, "value": on_hand * float(product["unit_cost"]),
            "outbound": used, "turnover": turnover, "coverage": coverage, "abc": abc.get(sku, "C"), "xyz": xyz,
            "reorder_point": reorder_point, "incoming": incoming, "expected": expected, "suggested_order": suggested,
            "last_movement": last_movement.get(sku, ""), "days_since_move": days_since_move,
            "status": status, "severity": severity,
        })
    rows.sort(key=lambda item: (-item["severity"], -item["value"], item["sku"]))
    total_value = sum(item["value"] for item in rows)
    avg_inventory_value = sum(
        (statistics.fmean(snapshot_qty[sku]) if snapshot_qty[sku] else current.get(sku, (0, 0))[0]) * float(product["unit_cost"])
        for sku, product in products.items()
    )
    turnover = total_outbound_value / avg_inventory_value if avg_inventory_value > 0 else None
    return {
        "mode": "standard",
        "period_days": period_days,
        "latest_snapshot": latest_date,
        "uploads": uploads,
        "rows": rows,
        "kpis": {
            "stock_value": total_value,
            "turnover": turnover,
            "days_inventory": (period_days / turnover) if turnover and turnover > 0 else None,
            "risk_count": sum(1 for row in rows if row["severity"] >= 3),
            "dead_value": sum(row["value"] for row in rows if row["status"] == "Nem forgó készlet"),
            "sku_count": len(rows),
        },
    }


def import_erp_exports(files: dict[str, tuple[str, bytes]], dataset: str = "parts") -> dict:
    """Import the native ERP stock overview and consumption reports."""
    stock_upload = files.get("stock_file")
    consumption_uploads = [
        (window_days, files[field_name])
        for field_name, window_days in (
            ("consumption_365_file", 365),
            ("consumption_90_file", 90),
            ("consumption_7_file", 7),
        )
        if files.get(field_name)
    ]
    if stock_upload is None and not consumption_uploads:
        raise ValueError("Legalább egy készletáttekintést vagy fogyási riportot válassz ki.")

    tables = _erp_table_names(dataset)
    results = []
    if stock_upload is not None:
        file_name, payload = stock_upload
        rows = _read_erp_stock(file_name, payload)
        upload_id = hashlib.sha256(payload).hexdigest()
        uploaded_at = datetime.now().isoformat(timespec="seconds")
        with _connect() as conn:
            duplicate = conn.execute(f"SELECT 1 FROM {tables['reports']} WHERE upload_id=?", (upload_id,)).fetchone()
            # The stock report is a replaceable current-state snapshot. Reparse
            # even an identical file so importer improvements take effect.
            conn.execute(f"DELETE FROM {tables['stock']}")
            conn.executemany(
                f"""
                INSERT INTO {tables['stock']} VALUES
                (:sku,:category,:description,:physical_stock,:available_stock,:open_demand,
                 :confirmed_demand,:open_purchase,:safety_stock,:storage_capacity,:length_value,
                 :width_value,:color,:inventory_flag,:uploaded_at,:upload_id)
                """,
                [dict(row, uploaded_at=uploaded_at, upload_id=upload_id) for row in rows],
            )
            if not duplicate:
                conn.execute(
                    f"INSERT INTO {tables['reports']} VALUES (?,?,?,?,?,?)",
                    (upload_id, uploaded_at, Path(file_name).name, "stock", 0, len(rows)),
                )
            conn.commit()
            results.append({"type": "Készletáttekintés", "rows": len(rows), "duplicate": bool(duplicate)})

    for expected_window_days, (file_name, payload) in consumption_uploads:
        window_days, rows = _read_erp_consumption(file_name, payload, expected_window_days)
        upload_id = hashlib.sha256(payload).hexdigest()
        uploaded_at = datetime.now().isoformat(timespec="seconds")
        with _connect() as conn:
            duplicate = conn.execute(f"SELECT 1 FROM {tables['reports']} WHERE upload_id=?", (upload_id,)).fetchone()
            if not duplicate:
                conn.executemany(
                    f"""
                    INSERT INTO {tables['consumption']} VALUES
                    (:upload_id,:uploaded_at,:window_days,:sku,:category,:description,:monthly_average,
                     :report_stock,:unit,:safety_stock,:inventory_flag,:long_description)
                    """,
                    [
                        dict(row, upload_id=upload_id, uploaded_at=uploaded_at, window_days=window_days)
                        for row in rows
                    ],
                )
                conn.execute(
                    f"INSERT INTO {tables['reports']} VALUES (?,?,?,?,?,?)",
                    (upload_id, uploaded_at, Path(file_name).name, "consumption", window_days, len(rows)),
                )
                conn.commit()
            results.append({"type": f"{window_days} napos fogyás", "rows": len(rows), "duplicate": bool(duplicate)})
    return {"reports": results}


def erp_dashboard(period_days: int = 90, dataset: str = "parts") -> dict:
    """Build the stock-level and turnover dashboard from the native ERP exports."""
    selected_window = 365 if period_days == 365 else (90 if period_days in {90, 180} else 7)
    tables = _erp_table_names(dataset)
    with _connect() as conn:
        stock = {row["sku"]: dict(row) for row in conn.execute(f"SELECT * FROM {tables['stock']}")}
        excluded_skus = {row["sku"] for row in conn.execute(f"SELECT sku FROM {tables['exclusions']}")}
        discontinued_skus = {
            row["sku"] for row in conn.execute(f"SELECT sku FROM {tables['settings']} WHERE discontinued=1")
        }
        reports = [dict(row) for row in conn.execute(f"SELECT * FROM {tables['reports']} ORDER BY uploaded_at DESC LIMIT 12")]
        latest_by_window: dict[int, dict[str, dict]] = {}
        for window in (7, 90, 365):
            report = conn.execute(
                f"SELECT upload_id, uploaded_at FROM {tables['reports']} WHERE report_type='consumption' AND window_days=? ORDER BY uploaded_at DESC LIMIT 1",
                (window,),
            ).fetchone()
            if report:
                latest_by_window[window] = {
                    row["sku"]: dict(row)
                    for row in conn.execute(f"SELECT * FROM {tables['consumption']} WHERE upload_id=?", (report["upload_id"],))
                }

    consumption_7 = latest_by_window.get(7, {})
    consumption_90 = latest_by_window.get(90, {})
    consumption_365 = latest_by_window.get(365, {})
    selected = latest_by_window.get(selected_window) or consumption_90 or consumption_365 or consumption_7
    rows = []
    for sku, item in stock.items():
        if sku in excluded_skus:
            continue
        if dataset == "doors" and sku.upper().startswith("NFAY"):
            continue
        c7 = float(consumption_7.get(sku, {}).get("monthly_average", 0) or 0)
        c90 = float(consumption_90.get(sku, {}).get("monthly_average", 0) or 0)
        c365 = float(consumption_365.get(sku, {}).get("monthly_average", 0) or 0)
        if dataset == "doors" and c90 <= 0 and c365 <= 0:
            continue
        selected_avg = float(selected.get(sku, {}).get("monthly_average", 0) or 0)
        policy_monthly_average = c90 or c365 or c7
        physical = float(item["physical_stock"])
        available = float(item["available_stock"])
        open_demand = float(item["open_demand"])
        confirmed_demand = float(item["confirmed_demand"])
        # Do not let empty, inactive ERP master-data rows distort any KPI,
        # chart or list. Negative positions remain visible for reconciliation.
        no_stock = abs(physical) < 1e-9 and abs(available) < 1e-9
        no_consumption = all(value <= 0 for value in (c7, c90, c365))
        no_demand = abs(open_demand) < 1e-9 and abs(confirmed_demand) < 1e-9
        if (dataset == "doors" and no_stock and no_demand) or (dataset != "doors" and no_stock and no_consumption):
            continue
        incoming = float(item["open_purchase"])
        discontinued = sku in discontinued_skus
        safety_months = 1 if dataset == "doors" else 3
        capacity_months = 2 if dataset == "doors" else 6
        safety = policy_monthly_average * safety_months
        capacity = policy_monthly_average * capacity_months
        projected = available - open_demand + incoming
        coverage_months = available / selected_avg if selected_avg > 0 else None
        months_to_stockout = max(0.0, projected) / policy_monthly_average if policy_monthly_average > 0 else None
        turnover = selected_avg * 12 / physical if physical > 0 else None
        turnover_90 = c90 * 12 / physical if physical > 0 and c90 > 0 else None
        turnover_365 = c365 * 12 / physical if physical > 0 and c365 > 0 else None
        turnover_months_90 = physical / c90 if physical > 0 and c90 > 0 else None
        turnover_months_365 = physical / c365 if physical > 0 and c365 > 0 else None
        trend_90_365 = c90 / c365 - 1 if c365 > 0 else None
        trend_7_90 = c7 / c90 - 1 if c90 > 0 and c7 > 0 else None
        capacity_use = physical / capacity if capacity > 0 else None
        if available <= 0 and policy_monthly_average > 0:
            base_status, base_severity = "Készlethiány", 5
        elif projected < safety and policy_monthly_average > 0:
            base_status, base_severity = "Beérkezővel is kritikus", 4
        elif available < safety and policy_monthly_average > 0:
            base_status, base_severity = "Biztonsági szint alatt", 3
        elif c90 == 0 and c365 == 0 and physical > 0:
            base_status, base_severity = "365 napja nem forgó", 2
        elif capacity_use is not None and capacity_use > 1:
            base_status, base_severity = "Tárolható mennyiség felett", 2
        else:
            base_status, base_severity = "Rendben", 0
        status = "Kifutó" if discontinued else base_status
        severity = 0 if discontinued else base_severity
        target_stock = max(safety, selected_avg * 2)
        base_suggested_order = max(0.0, target_stock - projected)
        suggested_order = 0.0 if discontinued else base_suggested_order
        rows.append({
            "sku": sku,
            "name": item["description"],
            "category": _canonical_category(
                selected.get(sku, {}).get("category")
                or consumption_365.get(sku, {}).get("category")
                or item["category"]
            ),
            "unit": (selected.get(sku, {}).get("unit") or "ME"),
            "physical": physical,
            "available": available,
            "open_demand": open_demand,
            "confirmed_demand": confirmed_demand,
            "incoming": incoming,
            "projected": projected,
            "safety": safety,
            "capacity": capacity,
            "capacity_use": capacity_use,
            "monthly_7": c7,
            "monthly_90": c90,
            "monthly_365": c365,
            "selected_average": selected_avg,
            "policy_monthly_average": policy_monthly_average,
            "coverage_months": coverage_months,
            "months_to_stockout": months_to_stockout,
            "turnover": turnover,
            "turnover_90": turnover_90,
            "turnover_365": turnover_365,
            "turnover_months_90": turnover_months_90,
            "turnover_months_365": turnover_months_365,
            "trend_90_365": trend_90_365,
            "trend_7_90": trend_7_90,
            "suggested_order": suggested_order,
            "base_suggested_order": base_suggested_order,
            "status": status,
            "severity": severity,
            "base_status": base_status,
            "base_severity": base_severity,
            "discontinued": discontinued,
            "erp_reconcile": available < 0 or physical < 0,
        })
    rows.sort(
        key=lambda row: (
            bool(row["discontinued"]),
            row["months_to_stockout"] is None,
            row["months_to_stockout"] if row["months_to_stockout"] is not None else math.inf,
            -row["severity"],
            row["sku"],
        )
    )
    statistics_rows = [row for row in rows if not row["discontinued"]]
    moving = [row for row in statistics_rows if row["selected_average"] > 0]
    turnover_values = [row["turnover"] for row in moving if row["turnover"] is not None]
    turnover_90_values = [row["turnover_90"] for row in statistics_rows if row["turnover_90"] is not None]
    turnover_365_values = [row["turnover_365"] for row in statistics_rows if row["turnover_365"] is not None]
    turnover_months_90_values = [row["turnover_months_90"] for row in statistics_rows if row["turnover_months_90"] is not None]
    turnover_months_365_values = [row["turnover_months_365"] for row in statistics_rows if row["turnover_months_365"] is not None]
    coverage_values = [row["coverage_months"] for row in moving if row["coverage_months"] is not None]
    stock_report = next((report for report in reports if report["report_type"] == "stock"), None)
    return {
        "mode": "erp",
        "dataset": dataset,
        "period_days": period_days,
        "selected_window": selected_window if latest_by_window.get(selected_window) else (90 if consumption_90 else 365),
        "latest_snapshot": stock_report["uploaded_at"] if stock_report else "",
        "reports": reports,
        "available_windows": sorted(latest_by_window),
        "rows": rows,
        "kpis": {
            "sku_count": len(statistics_rows),
            "moving_count": len(moving),
            "shortage_count": sum(row["severity"] >= 3 for row in statistics_rows),
            "negative_count": sum(row["available"] < 0 for row in statistics_rows),
            "erp_reconcile_count": sum(row["available"] < 0 or row["physical"] < 0 for row in statistics_rows),
            "dead_count": sum(row["status"] == "365 napja nem forgó" for row in statistics_rows),
            "overstock_count": sum(row["capacity_use"] is not None and row["capacity_use"] > 1 for row in statistics_rows),
            "median_turnover": statistics.median(turnover_values) if turnover_values else None,
            "median_turnover_90": statistics.median(turnover_90_values) if turnover_90_values else None,
            "median_turnover_365": statistics.median(turnover_365_values) if turnover_365_values else None,
            "median_turnover_months_90": statistics.median(turnover_months_90_values) if turnover_months_90_values else None,
            "median_turnover_months_365": statistics.median(turnover_months_365_values) if turnover_months_365_values else None,
            "median_coverage": statistics.median(coverage_values) if coverage_values else None,
            "incoming_count": sum(row["incoming"] > 0 for row in statistics_rows),
            "discontinued_count": sum(row["discontinued"] for row in rows),
        },
    }


def _read_erp_stock(file_name: str, payload: bytes) -> list[dict]:
    worksheet = _load_erp_sheet(file_name, payload)
    headers = [_text(value) for value in next(worksheet.iter_rows(values_only=True))]
    required = {
        "sku": "Alkatr.-szám",
        "description": "Alkatr.-leírás",
        "physical_stock": "Rend.áll.rakt.készl. ME",
        "available_stock": "Rend.áll",
        "open_demand": "Nyitott igény",
        "confirmed_demand": "Visszaigazolt igény",
        "open_purchase": "Nyitott beszerzés",
        "safety_stock": "Biztonsági készlet",
        "storage_capacity": "Tárolh.menny.",
    }
    columns = _erp_columns(headers, required, "készletáttekintés")
    optional = {name: headers.index(header) if header in headers else None for name, header in {
        "length_value": "HOSSZ", "width_value": "SZELESSEG", "color": "SZIN_NETTFR.Desc", "inventory_flag": "Leltarbol_ki",
    }.items()}
    rows, category = [], ""
    for values in worksheet.iter_rows(min_row=2, values_only=True):
        first, second = _text(values[0] if values else ""), _text(values[1] if len(values) > 1 else "")
        if first == "-" or second.lower().startswith(("icg kód :", "icg név :")):
            category = second.split(":", 1)[-1].strip()
            continue
        sku = _text(values[columns["sku"]] if columns["sku"] < len(values) else "")
        if not sku:
            continue
        def num(key: str) -> float:
            value = values[columns[key]] if columns[key] < len(values) else None
            return _erp_number(value)
        physical_value = values[columns["physical_stock"]] if columns["physical_stock"] < len(values) else None
        available_value = values[columns["available_stock"]] if columns["available_stock"] < len(values) else None
        physical = (
            _erp_number(available_value)
            if physical_value in (None, "") and available_value not in (None, "")
            else _erp_number(physical_value)
        )
        rows.append({
            "sku": sku, "category": category or "Nincs kategória",
            "description": _text(values[columns["description"]]) or sku,
            "physical_stock": physical,
            "available_stock": physical if available_value in (None, "") else _erp_number(available_value),
            "open_demand": num("open_demand"), "confirmed_demand": num("confirmed_demand"),
            "open_purchase": num("open_purchase"), "safety_stock": num("safety_stock"),
            "storage_capacity": num("storage_capacity"),
            **{
                key: _text(values[index]) if index is not None and index < len(values) else ""
                for key, index in optional.items()
            },
        })
    if not rows:
        raise ValueError("A készletáttekintés nem tartalmaz feldolgozható cikksort.")
    return rows


def _read_erp_consumption(
    file_name: str,
    payload: bytes,
    expected_window_days: int | None = None,
) -> tuple[int, list[dict]]:
    match = re.search(r"(?<!\d)(7|90|365)\s*nap", file_name.lower())
    if match and expected_window_days is not None and int(match.group(1)) != expected_window_days:
        raise ValueError(
            f"A kiválasztott mező {expected_window_days} napos riportot vár, "
            f"de a fájlnév {match.group(1)} napos időszakot jelez: {file_name}"
        )
    if expected_window_days is not None:
        window_days = expected_window_days
    elif match:
        window_days = int(match.group(1))
    else:
        raise ValueError(f"A fogyási riport fájlnevében szerepeljen a 7, 90 vagy 365 napos időszak: {file_name}")
    worksheet = _load_erp_sheet(file_name, payload)
    headers = [_text(value) for value in next(worksheet.iter_rows(values_only=True))]
    required = {
        "sku": "Alkatr.-szám", "description": "Alkatr.-leírás", "monthly_average": "Havi_atlag",
        "report_stock": "Rend.áll.rakt.készl.", "unit": "Elsődleges ME",
        "safety_stock": "Biztonsági készlet",
    }
    columns = _erp_columns(headers, required, "fogyási riport")
    inventory_index = headers.index("Leltarbol_ki") if "Leltarbol_ki" in headers else None
    long_index = headers.index("LEIRAS") if "LEIRAS" in headers else None
    rows, category = [], ""
    for values in worksheet.iter_rows(min_row=2, values_only=True):
        first, second = _text(values[0] if values else ""), _text(values[1] if len(values) > 1 else "")
        if first == "-" or second.lower().startswith(("icg kód :", "icg név :")):
            category = second.split(":", 1)[-1].strip()
            continue
        sku = _text(values[columns["sku"]] if columns["sku"] < len(values) else "")
        if not sku:
            continue
        rows.append({
            "sku": sku, "category": category or "Nincs kategória",
            "description": _text(values[columns["description"]]) or sku,
            "monthly_average": _erp_number(values[columns["monthly_average"]]),
            "report_stock": _erp_number(values[columns["report_stock"]]),
            "unit": _text(values[columns["unit"]]) or "ME",
            "safety_stock": _erp_number(values[columns["safety_stock"]]),
            "inventory_flag": _text(values[inventory_index]) if inventory_index is not None else "",
            "long_description": _text(values[long_index]) if long_index is not None else "",
        })
    if not rows:
        raise ValueError("A fogyási riport nem tartalmaz feldolgozható cikksort.")
    return window_days, rows


def _load_erp_sheet(file_name: str, payload: bytes):
    if not file_name.lower().endswith((".xlsx", ".xlsm")):
        raise ValueError("Az ERP-export csak XLSX vagy XLSM lehet.")
    if len(payload) > 30 * 1024 * 1024:
        raise ValueError("Egy ERP-export legfeljebb 30 MB lehet.")
    return load_workbook(io.BytesIO(payload), read_only=True, data_only=True).active


def _erp_columns(headers: list[str], required: dict[str, str], report_name: str) -> dict[str, int]:
    # The ERP exports use a visually merged two-cell heading for the item number;
    # the actual item code is in the otherwise unnamed second column.
    columns = {}
    missing = []
    for key, header in required.items():
        if key == "sku" and header in headers:
            header_index = headers.index(header)
            next_index = header_index + 1
            # Some ERP variants use a merged heading and place the item code
            # in an unnamed next column; others store it directly below the
            # visible "Alkatr.-szám" heading.
            columns[key] = (
                next_index
                if next_index < len(headers) and not headers[next_index]
                else header_index
            )
        elif header in headers:
            columns[key] = headers.index(header)
        else:
            missing.append(header)
    if missing:
        raise ValueError(f"A(z) {report_name} hiányzó oszlopa(i): {', '.join(missing)}.")
    return columns


def _erp_number(value) -> float:
    if value in (None, ""):
        return 0.0
    try:
        return float(str(value).replace(" ", "").replace(",", "."))
    except (TypeError, ValueError):
        return 0.0


def _read_sheet(sheet, required_headers: tuple[str, ...], sheet_name: str) -> list[dict]:
    rows = sheet.iter_rows(values_only=True)
    header_row = next(rows, None)
    if not header_row:
        return []
    headers = [_text(value) for value in header_row]
    missing = [header for header in required_headers if header not in headers]
    if missing:
        raise ValueError(f"A „{sheet_name}” lapon hiányzó oszlop(ok): {', '.join(missing)}.")
    index = {header: headers.index(header) for header in required_headers}
    output = []
    for values in rows:
        if not any(value not in (None, "") for value in values):
            continue
        output.append({header: values[column] if column < len(values) else None for header, column in index.items()})
    return output


def _text(value) -> str:
    return str(value or "").strip()


def _required_text(row: dict, key: str, sheet: str, row_number: int) -> str:
    value = _text(row.get(key))
    if not value:
        raise ValueError(f"Hiányzó „{key}” érték: {sheet} {row_number}. sor.")
    return value


def _number(value, sheet: str, row_number: int, column: str, minimum: float | None = None) -> float:
    try:
        result = float(str(value).replace(" ", "").replace(",", "."))
    except (TypeError, ValueError):
        raise ValueError(f"Hibás szám a(z) „{column}” oszlopban: {sheet} {row_number}. sor.") from None
    if not math.isfinite(result) or (minimum is not None and result < minimum):
        raise ValueError(f"Érvénytelen érték a(z) „{column}” oszlopban: {sheet} {row_number}. sor.")
    return result


def _date_text(value, sheet: str, row_number: int) -> str:
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    text = _text(value)
    for fmt in ("%Y-%m-%d", "%Y.%m.%d", "%Y.%m.%d.", "%d.%m.%Y", "%d.%m.%Y."):
        try:
            return datetime.strptime(text, fmt).date().isoformat()
        except ValueError:
            pass
    raise ValueError(f"Hibás dátum: {sheet} {row_number}. sor. Használd az ÉÉÉÉ-HH-NN formátumot.")


def _movement_type(value) -> str:
    text = unicodedata.normalize("NFKC", _text(value)).upper().replace(" ", "_")
    aliases = {
        "ÉRTÉKESÍTÉS": "ELADAS", "ELADÁS": "ELADAS", "ELADAS": "ELADAS",
        "TERMELÉSI_FELHASZNÁLÁS": "TERMELÉSI_FELHASZNALAS",
        "ANYAGKIADÁS": "ANYAGKIADAS", "FELHASZNÁLÁS": "FELHASZNALAS",
    }
    return aliases.get(text, text)
