"""Shared parser for frozen stock quantity exports."""

from __future__ import annotations

import csv
import io
import re
import unicodedata
from pathlib import Path

from tools.excel import normalize_excel_payload

try:
    from openpyxl import load_workbook
except Exception:  # pragma: no cover
    load_workbook = None


def read_frozen_stock_quantities(file_name: str, payload: bytes) -> dict[str, str]:
    """Read part-number keyed quantities from the standard three-column export."""
    rows = _read_rows(file_name, payload)
    if not rows:
        raise ValueError("A feltöltött befagyasztott készletlista üres.")

    headers = [_fold(value) for value in rows[0]]
    part_index = _find_header(headers, (("alkatr", "szam"), ("cikk", "szam")))
    qty_index = _find_header(
        headers,
        (
            ("befagyott", "keszlet", "menny"),
            ("befagyasztott", "keszlet", "menny"),
            ("rend", "all", "rakt", "keszl"),
            ("konyvelesi", "mennyiseg"),
            ("raktari", "keszlet"),
        ),
    )
    missing = []
    if part_index is None:
        missing.append("Alkatr.-szám")
    if qty_index is None:
        missing.append("Befagyott készlet menny.")
    if missing:
        raise ValueError("Hiányzó kötelező oszlop: " + ", ".join(missing))

    quantities: dict[str, str] = {}
    display_parts: dict[str, str] = {}
    for row_number, row in enumerate(rows[1:], start=2):
        if part_index >= len(row):
            continue
        part_number = _clean(row[part_index])
        if not part_number:
            continue
        raw_quantity = row[qty_index] if qty_index < len(row) else ""
        quantity = _quantity_text(raw_quantity)
        if quantity is None:
            raise ValueError(f"Hibás befagyasztott készlet a(z) {row_number}. sorban: {part_number}.")
        key = part_number.casefold()
        if key in quantities and quantities[key] != quantity:
            raise ValueError(f"Az alkatrészszám eltérő darabszámmal többször szerepel: {part_number}.")
        quantities[key] = quantity
        display_parts[key] = part_number

    if not quantities:
        raise ValueError("A feltöltött listában nem találtam alkatrészszámot és darabszámot.")
    return quantities


def _read_rows(file_name: str, payload: bytes) -> list[list | tuple]:
    suffix = Path(file_name or "").suffix.lower()
    if suffix == ".csv":
        text = payload.decode("utf-8-sig", errors="replace")
        dialect = csv.Sniffer().sniff(text[:2048], delimiters=";,\t,")
        return list(csv.reader(io.StringIO(text), dialect))
    if load_workbook is None:
        raise RuntimeError("Az Excel olvasásához hiányzik az openpyxl csomag.")
    workbook = load_workbook(io.BytesIO(normalize_excel_payload(payload)), read_only=True, data_only=True)
    return list(workbook.active.iter_rows(values_only=True))


def _find_header(headers: list[str], alternatives: tuple[tuple[str, ...], ...]) -> int | None:
    for index, header in enumerate(headers):
        if any(all(token in header for token in tokens) for tokens in alternatives):
            return index
    return None


def _fold(value: object) -> str:
    text = unicodedata.normalize("NFKD", _clean(value).lower())
    text = "".join(char for char in text if not unicodedata.combining(char))
    return re.sub(r"[^a-z0-9]+", " ", text).strip()


def _clean(value: object) -> str:
    if value is None:
        return ""
    return re.sub(r"\s+", " ", str(value)).strip()


def _quantity_text(value: object) -> str | None:
    if value is None or (isinstance(value, str) and not value.strip()):
        return None
    text = _clean(value).replace(" ", "").replace(",", ".")
    try:
        number = float(text)
    except ValueError:
        return None
    if number.is_integer():
        return str(int(number))
    return f"{number:.3f}".rstrip("0").rstrip(".")
